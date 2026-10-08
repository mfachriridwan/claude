"""Word (.docx) versions of the research documents with native, editable Word equations (OMML).

The PDF generators are re-used unchanged: their layout helpers (P, table, figure, eqrow, ...) are replaced by versions
that emit Pandoc Markdown, every equation is passed as LaTeX ($$...$$), and pandoc converts the result to .docx, where
each equation is a Word equation object (Insert > Equation) that can be clicked, edited and checked.
Run after the PDFs' inputs exist:  python docs/make_word_versions.py
Requires pandoc (e.g. pip install pypandoc_binary) and python-docx.
"""
import html
import re
import subprocess
import sys
from pathlib import Path

DOCS = Path(__file__).resolve().parent
sys.path.insert(0, str(DOCS))
import doccommon                                    # noqa: E402
import reportlab.platypus as platypus               # noqa: E402

try:
    import pypandoc
    PANDOC = pypandoc.get_pandoc_path()
except Exception:                                   # fall back to a pandoc on PATH
    PANDOC = "pandoc"

TARGETS = ["make_manuscript_pdf.py", "make_methodology_pdf.py", "make_math_model_pdf.py"]

# ------------------------------------------------------------------ reportlab mini-HTML -> Pandoc Markdown
MD_SPECIAL = re.compile(r"([\\*_~^\[\]#|<>`$])")
TAG = re.compile(r"(<[^>]+>)")


def md_inline(t, in_table=False):
    out, stack = [], []
    for part in TAG.split(str(t)):
        if not part:
            continue
        if part.startswith("<"):
            tag = part.strip("<>/ ").split()[0].lower() if part.strip("<>/ ") else ""
            closing = part.startswith("</")
            if tag == "b": out.append("**")
            elif tag == "i": out.append("*")
            elif tag == "sub": out.append("~"); stack.append("sub") if not closing else stack.pop() if stack else None
            elif tag in ("super", "sup"): out.append("^"); stack.append("sup") if not closing else stack.pop() if stack else None
            elif tag == "br": out.append(" " if in_table else "\\\n")
            continue                                # font, a, span, ... dropped (text kept)
        txt = MD_SPECIAL.sub(r"\\\1", html.unescape(part).replace(" ", " "))
        if stack:                                   # spaces inside sub/superscript must be escaped
            txt = txt.replace(" ", "\\ ")
        out.append(txt)
    s = "".join(out)
    s = re.sub(r"\*\*\s*\*\*", "", s)
    return s.replace("\n", " ") if in_table else s


class Block(str):
    """A Markdown block; `kind` lets list items be joined into one list."""
    def __new__(cls, text, kind="para"):
        o = str.__new__(cls, text); o.kind = kind; return o


def P(t, st="body"):
    s = md_inline(t)
    if st == "title": return Block(f'::: {{custom-style="Title"}}\n{s}\n:::')
    if st == "subtitle": return Block(f'::: {{custom-style="Subtitle"}}\n{s}\n:::')
    if st == "h1": return Block(f"# {s}")
    if st == "h2": return Block(f"## {s}")
    if st == "cap": return Block(f'::: {{custom-style="Caption"}}\n{s}\n:::')
    if st == "ref": return Block(f'::: {{custom-style="Bibliography"}}\n{s}\n:::')
    return Block(s)


def bullets(items, st="body"):
    return [Block("- " + md_inline(t), "li") for t in items]


def table(rows, widths, header=1, zebra=True, font=None):
    n = max(len(r) for r in rows)
    cells = [[md_inline(c, True) if isinstance(c, str) else md_inline(str(c), True) for c in r] + [""] * (n - len(r)) for r in rows]
    tot = sum(widths[:n]) or 1
    dashes = ["-" * max(3, int(round(60 * w / tot))) for w in widths[:n]] + ["---"] * (n - len(widths))
    lines = ["| " + " | ".join(c or " " for c in cells[0]) + " |", "|" + "|".join(dashes) + "|"]
    lines += ["| " + " | ".join(c or " " for c in r) + " |" for r in cells[1:]]
    return Block("\n".join(lines), "table")


def flatten(x):
    if isinstance(x, (list, tuple)):
        out = []
        for y in x: out += flatten(y)
        return out
    return [x] if x not in (None, "") else []


def boxed(flow, width=17.0, bg=None):
    inner = join_blocks(flatten(flow))
    return Block(f'::: {{custom-style="Block Text"}}\n{inner}\n:::')


def img(path, width_cm):
    return Block(f"![]({Path(path).resolve()}){{width={min(width_cm, 16.5):.1f}cm}}")


def figure(path, width_cm, caption):
    return Block(f"![{md_inline(caption)}]({Path(path).resolve()}){{width={min(width_cm, 16.5):.1f}cm}}")


def eq(tex, name=None, width_cm=None, fs=13):
    return Block(f"$${tex}$$")


def eqrow(tex, name=None, label="", fs=13):
    return Block(f"$${tex}\\qquad({label})$$")


def page_deco(title):
    return lambda canvas, doc: None


PAGEBREAK = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```'


def join_blocks(blocks):
    out, prev = [], None
    for b in blocks:
        kind = getattr(b, "kind", "para")
        if out and kind == "li" and prev == "li":
            out[-1] += "\n" + b
        else:
            out.append(str(b))
        prev = kind
    return "\n\n".join(o for o in out if o.strip())


class FakeDoc:
    def __init__(self, path, **kw):
        self.path = Path(str(path)); self.title = kw.get("title", ""); self.author = kw.get("author", "")

    def build(self, story, **kw):
        md = join_blocks(flatten(story))
        meta = f"---\ntitle-meta: \"{self.title}\"\nauthor-meta: \"{self.author}\"\nlang: en-GB\n---\n\n"
        out = self.path.with_suffix(".docx")
        src = DOCS / "word_build" / (self.path.stem + ".md"); src.parent.mkdir(exist_ok=True)
        src.write_text(meta + md, encoding="utf-8")
        ref = reference_doc()
        r = subprocess.run([PANDOC, str(src), "-f", "markdown+tex_math_dollars+subscript+superscript+pipe_tables",
                            "-o", str(out), "--reference-doc", str(ref), "--columns=40"], capture_output=True, text=True)
        if r.returncode:
            raise RuntimeError(r.stderr)
        warn = [l for l in r.stderr.splitlines() if "math" in l.lower() or "tex" in l.lower()]
        postprocess(out)
        print("wrote", out, f"({len(warn)} math warnings)" if warn else "")
        for w in warn[:10]: print("   ", w)


# ------------------------------------------------------------------ reference document and post-processing
def reference_doc():
    ref = DOCS / "word_build" / "reference.docx"
    if ref.exists():
        return ref
    ref.parent.mkdir(exist_ok=True)
    data = subprocess.run([PANDOC, "--print-default-data-file", "reference.docx"], capture_output=True).stdout
    ref.write_bytes(data)
    import docx
    from docx.shared import Pt, RGBColor, Cm
    d = docx.Document(str(ref))
    for s in d.styles:
        try:
            if s.font is not None:
                s.font.name = "Times New Roman"
                rpr = s.element.get_or_add_rPr(); rf = rpr.find(docx.oxml.ns.qn("w:rFonts"))
                if rf is None:
                    rf = docx.oxml.OxmlElement("w:rFonts"); rpr.append(rf)
                for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(docx.oxml.ns.qn(a), "Times New Roman")
                for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                    if rf.get(docx.oxml.ns.qn(a)) is not None: del rf.attrib[docx.oxml.ns.qn(a)]
        except Exception:
            pass
    sizes = {"Normal": 11, "Body Text": 11, "First Paragraph": 11, "Title": 17, "Subtitle": 11.5, "Heading 1": 13.5,
             "Heading 2": 11.5, "Caption": 9.5, "Image Caption": 9.5, "Table Caption": 9.5, "Bibliography": 9.5,
             "Compact": 9.5}
    for name, sz in sizes.items():
        try:
            st = d.styles[name]; st.font.size = Pt(sz); st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x36)
        except KeyError:
            pass
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.0)
    d.save(str(ref))
    return ref


def postprocess(path):
    """Borders, header shading and compact text for tables; justified body text."""
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = docx.Document(str(path))
    def insert_ordered(parent, el, after_tags):
        """Insert el after the last existing child whose tag is in after_tags (schema order), else first."""
        idx = -1
        for i, ch in enumerate(parent):
            if ch.tag in {qn(x) for x in after_tags}: idx = i
        parent.insert(idx + 1, el)

    TBL_BEFORE = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                  "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd")
    TC_BEFORE = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders")
    for t in d.tables:
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "A6A6A6"); b.append(e)
        insert_ordered(tblPr, b, TBL_BEFORE)
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
                    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "DCE6F0")
                    insert_ordered(tcPr, sh, TC_BEFORE)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8)
                        if i == 0: r.font.bold = True
    for pm in d.element.body.iter(qn("w:pgMar")):           # complete page margins (schema requires all attributes)
        for a, v in (("w:header", "708"), ("w:footer", "708"), ("w:gutter", "0")):
            if pm.get(qn(a)) is None: pm.set(qn(a), v)
    for p in d.paragraphs:
        if p.style.name in ("Body Text", "First Paragraph") and not p._p.xpath(".//m:oMath"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.save(str(path))


# ------------------------------------------------------------------ run the generators with the Markdown backend
def patch():
    for name, fn in (("P", P), ("bullets", bullets), ("table", table), ("boxed", boxed), ("img", img), ("figure", figure),
                     ("eq", eq), ("eqrow", eqrow), ("page_deco", page_deco)):
        setattr(doccommon, name, fn)
    platypus.SimpleDocTemplate = FakeDoc
    platypus.Spacer = lambda *a, **k: None
    platypus.PageBreak = lambda *a, **k: Block(PAGEBREAK)
    platypus.CondPageBreak = lambda *a, **k: None
    platypus.KeepTogether = lambda flow, *a, **k: flatten(flow)
    doccommon.Spacer, doccommon.KeepTogether = platypus.Spacer, platypus.KeepTogether


if __name__ == "__main__":
    patch()
    import os
    os.chdir(DOCS)
    for f in (sys.argv[1:] or TARGETS):
        ns = {"__name__": "__main__", "__file__": str(DOCS / f)}
        exec(compile((DOCS / f).read_text(), str(DOCS / f), "exec"), ns)
