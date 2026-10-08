"""Write the one-city input templates (XLSX + CSV), blank and with the Kota Padang example.
Run: python make_single_city_template.py"""
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from mswpath import MSWModel
from mswpath import single as S

OUT = Path("templates"); OUT.mkdir(exist_ok=True)
M = MSWModel("data")
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E78"); INP = PatternFill("solid", fgColor="FFF2CC"); EXF = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

GUIDE = [
    ("MSW recovery-pathway screening: one-city input template / Templat input satu kota", True),
    ("", False),
    ("HOW TO FILL IN / CARA MENGISI", True),
    ("1. Sheet 'city_data': type your values in the yellow column 'value'. Grey column = example (Kota Padang, RIPS).", False),
    ("   Isi kolom kuning 'value'. Kolom abu-abu = contoh (Kota Padang, RIPS).", False),
    ("2. Sheet 'composition': wet-mass % per RIPS category for the domestic stream (required) and the non-domestic stream", False),
    ("   (optional). Leave a cell BLANK if the category is not reported; blank is never read as zero in the data,", False),
    ("   but a not-reported category gets zero share in the model (a stated assumption). The TOTAL row should be about 100.", False),
    ("   Isi % berat basah per kategori. Kosongkan bila tidak dilaporkan (bukan nol). Baris TOTAL sebaiknya sekitar 100.", False),
    ("3. Sheet 'local_parameters' (optional): if you have a local value (e.g. measured moisture, landfill cost, RDF price),", False),
    ("   type it in 'your_central' (and optionally 'your_low', 'your_high', 'your_source'). Blank rows keep the default.", False),
    ("   If only your_central is given, the default relative range is kept around your value.", False),
    ("   Opsional: isi nilai lokal Anda; baris kosong memakai nilai bawaan model.", False),
    ("4. Save as .xlsx (or as three .csv files with the same columns) and upload it in MSW_Single_City_Colab.ipynb.", False),
    ("", False),
    ("RULES APPLIED BY THE MODEL / ATURAN MODEL", True),
    ("- Tonnage is projected once to 2025: Q2025 = Qt x (1 + g)^(2025 - t). Domestic and non-domestic use the same factor.", False),
    ("- Composition is normalised per stream and weighted by the 2025 tonnage of each stream.", False),
    ("- Either lat/lon of the planned site (nearest cement kiln is found) or the road distance to a kiln is needed.", False),
    ("- Options compared per tonne of mixed MSW at the facility gate in 2025: sanitary landfill + flare (SL), WtE (S1),", False),
    ("  RDF to cement kiln (S2), AD of food waste (S3), PHB from landfill gas (S4), RDF + AD (S5).", False),
    ("- Decision criterion: carbon-inclusive cost C + pG/1000 at fixed carbon values 0, 2, 25, 50, 100 USD/t CO2e.", False),
    ("- Screening level (ex ante): results show which options merit a feasibility study, not how a plant will perform.", False),
]


def style_header(ws, row, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row=row, column=j); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX


def write_frame(ws, df, input_cols=(), example_cols=(), widths=None):
    ws.append(list(df.columns)); style_header(ws, 1, len(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
    for i in range(2, ws.max_row + 1):
        for j, col in enumerate(df.columns, start=1):
            c = ws.cell(row=i, column=j); c.font = Font(name=F, size=10, color="0000FF" if col in input_cols else "000000")
            c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
            if col in input_cols: c.fill = INP
            elif col in example_cols: c.fill = EXF
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    ws.freeze_panes = "B2"


def build_xlsx(path, cd, comp, par):
    wb = Workbook()
    g = wb.active; g.title = "guide_panduan"
    for t, b in GUIDE:
        g.append([t]); g.cell(row=g.max_row, column=1).font = Font(name=F, bold=b, size=12 if b else 10)
    g.append([]); g.append(["Legend / Keterangan"]); g.cell(row=g.max_row, column=1).font = Font(name=F, bold=True)
    for txt, fill in (("yellow cell, blue text = your input / isian Anda", INP), ("grey cell = example (Kota Padang) / contoh", EXF)):
        g.append([txt]); g.cell(row=g.max_row, column=1).fill = fill; g.cell(row=g.max_row, column=1).font = Font(name=F, size=10)
    g.column_dimensions["A"].width = 125

    ws = wb.create_sheet("city_data")
    write_frame(ws, cd, input_cols=("value",), example_cols=("example_kota_padang",), widths=[26, 18, 9, 9, 22, 60, 60])
    rows = {f: i + 2 for i, f in enumerate(cd.field)}
    dv = DataValidation(type="list", formula1='"Jamali,Sumatera,Mahakam"', allow_blank=False); ws.add_data_validation(dv)
    dv.add(f"B{rows['grid_region']}")
    dv2 = DataValidation(type="list", formula1='"yes,no"', allow_blank=True); ws.add_data_validation(dv2)
    dv2.add(f"B{rows['organik_lumped']}")
    for fld in ("tonnage_year", "composition_year"):
        dv3 = DataValidation(type="whole", operator="between", formula1="1990", formula2="2025"); ws.add_data_validation(dv3)
        dv3.add(f"B{rows[fld]}")
    ws[f"B{rows['grid_region']}"].comment = Comment("Jamali = Java-Madura-Bali; Sumatera; Mahakam = East Kalimantan", "template")
    ws[f"B{rows['kiln_road_km']}"].comment = Comment("Leave blank to compute the distance from lat/lon to the nearest kiln "
                                                     "in data/cement_kilns.csv", "template")

    wc = wb.create_sheet("composition")
    write_frame(wc, comp, input_cols=("domestic_pct", "nondomestic_pct"),
                example_cols=("example_domestic_pct", "example_nondomestic_pct"), widths=[16, 22, 20, 15, 15, 17, 19, 22])
    n = wc.max_row
    wc.append(["TOTAL", "should be about 100", "harus sekitar 100", ""] +
              [f"=SUM({col}2:{col}{n})" for col in ("E", "F", "G", "H")])
    for j in range(1, 9):
        c = wc.cell(row=n + 1, column=j); c.font = Font(name=F, bold=True); c.border = BOX
        if j >= 5: c.number_format = "0.00"

    wp = wb.create_sheet("local_parameters")
    write_frame(wp, par, input_cols=("your_central", "your_low", "your_high", "your_source"),
                widths=[17, 15, 52, 15, 14, 12, 12, 13, 11, 11, 30])
    for i in range(2, wp.max_row + 1):
        for j in (5, 6, 7):
            wp.cell(row=i, column=j).number_format = "General"
    wb.calculation.fullCalcOnLoad = True          # TOTAL row is computed by Excel/Sheets when the file opens
    wb.save(path)


cd, comp, par = S.blank_template(M, S.example_padang("data"))
build_xlsx(OUT / "single_city_template.xlsx", cd, comp, par)
cd.to_csv(OUT / "single_city_city_data.csv", index=False)
comp.to_csv(OUT / "single_city_composition.csv", index=False)
par.to_csv(OUT / "single_city_local_parameters.csv", index=False)
cde, compe, pare = S.filled_example(M, "data")
build_xlsx(OUT / "single_city_example_kota_padang.xlsx", cde, compe, pare)
cde.to_csv(OUT / "example_kota_padang_city_data.csv", index=False)
compe.to_csv(OUT / "example_kota_padang_composition.csv", index=False)
pare.to_csv(OUT / "example_kota_padang_local_parameters.csv", index=False)
print("templates written:", sorted(p.name for p in OUT.glob("*single_city*")) + sorted(p.name for p in OUT.glob("example_kota*")))
